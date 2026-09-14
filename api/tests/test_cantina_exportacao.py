"""As exportações por tela da cantina (decisão de 14/09).

⚠️ O que estes testes trancam é quem vê o quê, não o layout do arquivo:

  * a restrição alimentar sai por extenso SÓ na exportação da cantina — na da
    coordenação ela é, no máximo, um "sim" (docs/38 §2.6);
  * a cantina exporta só os pedidos DELA — o `cantina_id` vem do token;
  * o CSV e o XLSX de pedidos saem do mesmo montador, e por isso têm o mesmo
    cabeçalho;
  * o período tem teto, porque não há paginação em lugar nenhum.
"""

import asyncio
import io
from datetime import date

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app import cantina_relatorio
from app.routes import cantina as rotas

from .fake_postgrest import FakeCliente
from .test_cantina_relatorio import _banco

DE = date(2026, 3, 10)
ATE = date(2026, 3, 20)


def _com_restricao() -> dict:
    db = _banco()
    db["aluno"]["A1"]["restricao_alimentar"] = "sem lactose"
    db["aluno"]["A1"]["matricula"] = "0001"
    db["aluno"]["A1"]["ativo"] = True
    db["aluno"]["A2"]["restricao_alimentar"] = None
    db["aluno"]["A2"]["matricula"] = "0002"
    db["aluno"]["A2"]["ativo"] = True
    db["direito_refeicao_aluno"] = {
        "d1": {"id": "d1", "aluno_id": "A1", "refeicao": "almoco"},
    }
    db["usuario_cantina"] = {
        "u1": {"id": "u1", "cantina_id": "ari", "email": "copa@ari.br", "nome": "Dona Maria",
               "ativo": True, "ultimo_login_em": None},
    }
    return db


def _celulas(conteudo: bytes) -> list[list]:
    ws = load_workbook(io.BytesIO(conteudo)).worksheets[0]
    return [[c.value for c in linha] for linha in ws.iter_rows()]


class TestARestricaoAlimentar:
    def test_a_planilha_da_coordenacao_nao_tem_a_coluna(self):
        linhas = _celulas(cantina_relatorio.planilha_de_pedidos(
            FakeCliente(_com_restricao()), de=DE, ate=ATE,
        ))
        assert "Restrição alimentar" not in linhas[0]
        assert not any("lactose" in str(c) for linha in linhas for c in linha)

    def test_a_planilha_da_cantina_traz_o_texto(self):
        linhas = _celulas(cantina_relatorio.planilha_de_pedidos(
            FakeCliente(_com_restricao()), de=DE, ate=ATE, com_restricao=True,
        ))
        coluna = linhas[0].index("Restrição alimentar")
        assert "sem lactose" in [linha[coluna] for linha in linhas[1:]]

    def test_alunos_com_direito_diz_so_sim_nunca_o_texto(self):
        linhas = _celulas(cantina_relatorio.planilha_de_direitos(FakeCliente(_com_restricao())))
        assert linhas[0][-1] == "Tem restrição"
        assert not any("lactose" in str(c) for linha in linhas for c in linha)
        ana = next(l for l in linhas if l[0] == "Ana Beatriz")
        assert ana[-1] == "sim"


class TestOCsv:
    def test_mesmo_cabecalho_do_xlsx(self):
        db = _com_restricao()
        xlsx = _celulas(cantina_relatorio.planilha_de_pedidos(FakeCliente(db), de=DE, ate=ATE))
        csv = cantina_relatorio.csv_de_pedidos(FakeCliente(db), de=DE, ate=ATE).decode("utf-8-sig")
        assert csv.splitlines()[0].split(";") == xlsx[0]

    def test_abre_no_excel_brasileiro(self):
        csv = cantina_relatorio.csv_de_pedidos(FakeCliente(_com_restricao()), de=DE, ate=ATE)
        assert csv.startswith("﻿".encode())
        assert "18,00" in csv.decode("utf-8-sig")

    def test_nome_comecando_com_igual_nao_vira_formula(self):
        """Um nome vindo do Canvas começando com `=` viraria célula executável
        na máquina de quem abriu o CSV."""
        db = _com_restricao()
        db["aluno"]["A1"]["nome"] = "=HYPERLINK(\"x\")"
        csv = cantina_relatorio.csv_de_pedidos(FakeCliente(db), de=DE, ate=ATE).decode("utf-8-sig")
        assert "'=HYPERLINK" in csv


class TestAsRotas:
    def test_a_cantina_so_exporta_os_pedidos_dela(self, monkeypatch):
        """A Food não pode baixar os pedidos da Cantina do Ari."""
        db = _com_restricao()
        monkeypatch.setattr(rotas, "get_supabase", lambda: FakeCliente(db))
        resposta = asyncio.run(rotas.pedidos_xlsx_da_cantina(DE, ATE, {"cantina_id": "food"}))
        linhas = _celulas(resposta.body)
        cantinas = {linha[1] for linha in linhas[1:]}
        assert cantinas == {"Food"}

    def test_periodo_invertido_e_recusado(self):
        with pytest.raises(HTTPException) as erro:
            rotas._periodo_valido(ATE, DE)
        assert erro.value.status_code == 422

    def test_periodo_acima_do_teto_e_recusado(self):
        with pytest.raises(HTTPException) as erro:
            rotas._periodo_valido(date(2026, 1, 1), date(2026, 12, 31))
        assert "93" in erro.value.detail

    def test_acesso_nao_tem_senha_nem_hash(self):
        linhas = _celulas(cantina_relatorio.planilha_de_acesso(FakeCliente(_com_restricao())))
        cabecalho = " ".join(str(c) for c in linhas[0]).lower()
        assert "senha" not in cabecalho and "hash" not in cabecalho
        assert any(linha[5] == "Dona Maria" for linha in linhas[1:])
