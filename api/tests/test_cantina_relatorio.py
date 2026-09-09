"""O relatório de custos em XLSX (docs/40 §12.11).

⚠️ O que estes testes trancam não é o layout da planilha — é o que ela AFIRMA:
que o custo de um período usa o preço CONGELADO no instante de cada pedido, e
não o valor de tabela de hoje. Sem isso, subir o almoço em abril mudaria o
relatório de março, e quem confere concluiria que o sistema errou.
"""

import io
from datetime import date

import pytest
from openpyxl import load_workbook

from app import cantina_relatorio

from .fake_postgrest import FakeCliente

DE = date(2026, 3, 10)
ATE = date(2026, 3, 20)


def _banco() -> dict:
    """Dois dias, duas cantinas, e o preço MUDANDO no meio do período."""
    return {
        "cantina": {
            "ari": {"id": "ari", "nome": "Cantina do Ari", "ativo": True},
            "food": {"id": "food", "nome": "Food", "ativo": True},
        },
        "cardapio": {
            "c1": {"id": "c1", "cantina_id": "ari", "data": "2026-03-12", "refeicao": "almoco"},
            "c2": {"id": "c2", "cantina_id": "food", "data": "2026-03-18", "refeicao": "almoco"},
        },
        "pedido_refeicao": {
            # 12/03, quando o almoço valia 18. 18/03, quando passou a valer 21.
            "p1": {"id": "p1", "cardapio_id": "c1", "aluno_id": "A1", "modo": "pedido",
                   "valor_cobrado": 18.0, "data": "2026-03-12", "refeicao": "almoco",
                   "criado_em": "2026-03-11T20:00:00Z", "retirado_em": None},
            "p2": {"id": "p2", "cardapio_id": "c1", "aluno_id": "A2", "modo": "pedido",
                   "valor_cobrado": 18.0, "data": "2026-03-12", "refeicao": "almoco",
                   "criado_em": "2026-03-11T20:05:00Z", "retirado_em": None},
            "p3": {"id": "p3", "cardapio_id": "c2", "aluno_id": "A1", "modo": "presencial",
                   "valor_cobrado": 21.0, "data": "2026-03-18", "refeicao": "almoco",
                   "criado_em": "2026-03-18T12:00:00Z", "retirado_em": "2026-03-18T12:07:00Z"},
        },
        "aluno": {
            "A1": {"id": "A1", "nome": "Ana Beatriz"},
            "A2": {"id": "A2", "nome": "Caio Rocha"},
        },
        "matricula_turma": {
            "m1": {"id": "m1", "aluno_id": "A1", "turma_id": "T1", "ativo_ate": None},
            "m2": {"id": "m2", "aluno_id": "A2", "turma_id": "T1", "ativo_ate": None},
        },
        "cardapio_bloco": {
            "b1": {"id": "b1", "cardapio_id": "c1", "nome": "Guarnição", "ordem": 1,
                   "escolhas_minimas": 0, "escolhas_maximas": 2,
                   "observacao": "a opção 4 anula a 1 e a 2"},
            "b2": {"id": "b2", "cardapio_id": "c1", "nome": "Proteína", "ordem": 2,
                   "escolhas_minimas": 1, "escolhas_maximas": 1, "observacao": None},
        },
        "cardapio_opcao": {
            "o1": {"id": "o1", "bloco_id": "b1", "nome": "Arroz", "ordem": 1},
            "o2": {"id": "o2", "bloco_id": "b1", "nome": "Feijão", "ordem": 2},
            "o3": {"id": "o3", "bloco_id": "b2", "nome": "Frango", "ordem": 1},
        },
        "pedido_refeicao_item": {
            "i1": {"id": "i1", "pedido_id": "p1", "opcao_id": "o1"},
            "i2": {"id": "i2", "pedido_id": "p1", "opcao_id": "o2"},
            "i3": {"id": "i3", "pedido_id": "p1", "opcao_id": "o3"},
        },
    }


@pytest.fixture
def planilha():
    conteudo = cantina_relatorio.montar(FakeCliente(_banco()), de=DE, ate=ATE)
    return load_workbook(io.BytesIO(conteudo))


def _linhas(ws) -> list[list]:
    return [[c.value for c in linha] for linha in ws.iter_rows()]


class TestAsAbas:
    def test_tem_as_abas_de_pedidos_custos_e_cardapio(self, planilha):
        assert "Pedidos" in planilha.sheetnames
        assert "Custos · por dia" in planilha.sheetnames
        assert "Custos · por turma" in planilha.sheetnames
        assert "Custos · por aluno" in planilha.sheetnames
        assert "Cardápio" in planilha.sheetnames

    def test_a_aba_por_cantina_aparece_porque_ha_DUAS(self, planilha):
        """Com uma cantina só ela não existe: uma aba de uma linha é ruído, e
        ensina a ignorar as abas do arquivo."""
        assert "Custos · por cantina" in planilha.sheetnames


class TestOPrecoCongelado:
    def test_o_total_usa_o_valor_de_CADA_pedido(self, planilha):
        """O teste que dá sentido à migration 0054.

        Dois pedidos a 18 e um a 21 somam 57. Se o relatório usasse o valor de
        tabela de hoje, os três sairiam ao mesmo preço — e março mudaria de
        custo toda vez que o preço subisse.
        """
        linhas = _linhas(planilha["Custos · por aluno"])
        total = {linha[0]: linha[2] for linha in linhas[1:]}
        assert total["Ana Beatriz"] == 39.0  # 18 (12/03) + 21 (18/03)
        assert total["Caio Rocha"] == 18.0

    def test_por_dia_separa_as_duas_datas(self, planilha):
        linhas = _linhas(planilha["Custos · por dia"])
        por_dia = {linha[0]: (linha[1], linha[2]) for linha in linhas[1:]}
        assert por_dia["2026-03-12 · Almoço"] == (2, 36.0)
        assert por_dia["2026-03-18 · Almoço"] == (1, 21.0)


class TestAAbaDePedidos:
    def test_uma_coluna_por_bloco(self, planilha):
        cabecalho = _linhas(planilha["Pedidos"])[0]
        assert "Guarnição" in cabecalho
        assert "Proteína" in cabecalho
        # A ordem é a da BANDEJA, não alfabética.
        assert cabecalho.index("Guarnição") < cabecalho.index("Proteína")

    def test_quem_pega_na_hora_sai_com_TRAÇO_e_nao_vazio(self, planilha):
        """Célula vazia numa planilha lê-se como "faltou o dado"; o traço diz
        que não havia dado a ter (docs/40 §10.1)."""
        linhas = _linhas(planilha["Pedidos"])
        cabecalho = linhas[0]
        presencial = next(l for l in linhas[1:] if l[cabecalho.index("Modo")] == "retirada na hora")
        assert presencial[cabecalho.index("Guarnição")] == "—"

    def test_a_cantina_aparece_em_cada_linha(self, planilha):
        """Com duas cantinas, a planilha sem esta coluna junta duas cozinhas na
        mesma contagem."""
        linhas = _linhas(planilha["Pedidos"])
        coluna = linhas[0].index("Cantina")
        assert {l[coluna] for l in linhas[1:]} == {"Cantina do Ari", "Food"}

    def test_a_restricao_alimentar_NAO_entra(self, planilha):
        """Quem exporta custo não precisa de dado de saúde de menor
        (docs/38 §2.6)."""
        cabecalho = [str(c) for c in _linhas(planilha["Pedidos"])[0]]
        assert not any("estri" in c for c in cabecalho)


class TestAGradeDoCardapio:
    def test_a_coluna_obs_junta_o_limite_com_o_texto_livre(self, planilha):
        linhas = _linhas(planilha["Cardápio"])
        obs = [linha[-1] for linha in linhas[1:] if linha[-1]]
        assert any("máximo 2 opções" in str(o) and "anula a 1 e a 2" in str(o) for o in obs)

    def test_os_dias_viram_COLUNAS(self, planilha):
        cabecalho = [str(c) for c in _linhas(planilha["Cardápio"])[0]]
        assert cabecalho[0] == "Bloco"
        assert any("2026-03-12" in c for c in cabecalho)
