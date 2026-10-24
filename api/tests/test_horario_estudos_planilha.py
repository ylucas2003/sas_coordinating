"""A planilha do horário de estudos (`app/horario_estudos_planilha.py`, docs/42).

O que este arquivo trava, em ordem de gravidade:

  1. **Texto do aluno nunca vira fórmula.** O rótulo é texto livre, e o openpyxl
     grava como FÓRMULA toda string que começa com `=`. Sem o cinto, "=1+1"
     abriria como `2` — e um arquivo repassado ao grupo da turma seria célula
     executável na máquina de quem abriu. O teste RELÊ o arquivo salvo, porque é
     o que o Excel lê, e não o objeto em memória.
  2. **Dois blocos nunca se mesclam sobre a mesma célula.** Hora quebrada (19:07)
     arredonda para a linha de 15 minutos, e dois blocos próximos podiam cair na
     mesma linha — o openpyxl aceita a mescla sobreposta e o Excel recusa abrir o
     arquivo ("encontramos um problema com algum conteúdo").
  3. **O horário só do próprio aluno.** A rota lê com o `aluno_id` do token.
  4. **As cores são as do site, nos dois temas.** `PALETAS` é lida de volta de
     `web/styles/paleta.css`: se alguém mexer na paleta e esquecer a planilha, falha aqui.
  5. **A grade tem a mesma faixa de horas do PDF.** A regra mora duas vezes
     (aqui e em `horarioEstudosFolha.ts`), e estes casos são os mesmos do teste de
     lá: se um mudar sozinho, o PDF e o Excel passam a mostrar recortes diferentes.

Rodar:  cd api && ./.venv/bin/python -m pytest tests/test_horario_estudos_planilha.py -q
"""

from __future__ import annotations

import asyncio
import io
import re
from datetime import date
from itertools import combinations
from pathlib import Path

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from app.horario_estudos_planilha import (
    PALETAS,
    TemaDaPlanilha,
    faixa_de_horas,
    planilha_do_horario,
)
from app.schemas.horario_estudos import BlocoDeEstudo

from .fake_postgrest import FakeCliente

HOJE = date(2026, 10, 24)


def _b(id_: str, dia: int, hora: str, duracao: int, rotulo: str = "x") -> BlocoDeEstudo:
    return BlocoDeEstudo(id=id_, diaSemana=dia, horaInicio=hora, duracaoMinutos=duracao, rotulo=rotulo)


def _abrir(blocos, nome="BENNY", tema="dia"):
    """O arquivo SALVO e relido: é o que o Excel lê, não o objeto em memória."""
    return load_workbook(io.BytesIO(planilha_do_horario(blocos, nome, HOJE, tema)))


def _cor(celula_ou_fonte) -> str:
    """O hex de 6 dígitos de uma célula (fundo) ou de uma fonte (letra).

    O openpyxl guarda ARGB ('001B3F8B'); o que importa é o RGB.
    """
    cor = getattr(celula_ou_fonte, "fill", None)
    rgb = cor.fgColor.rgb if cor is not None else celula_ou_fonte.color.rgb
    return str(rgb)[-6:]


def _mesclas(ws) -> list[str]:
    return sorted(str(r) for r in ws.merged_cells.ranges)


# ─── A faixa de horas — os MESMOS casos de horarioEstudosFolha.test.ts ────


@pytest.mark.parametrize(
    ("inicio", "duracao", "esperado"),
    [
        ("19:00", 90, (17, 23)),     # a folga se divide: não é 18–24
        ("00:00", 60, (0, 6)),       # bate no zero: o que sobrou vai para baixo
        ("22:00", 120, (18, 24)),    # bate na meia-noite: o que sobrou vai para cima
        ("23:00", 60, (18, 24)),
        ("08:00", 600, (8, 18)),     # já passa do mínimo: não cresce
        ("06:00", 1020, (6, 23)),
    ],
)
def test_faixa_de_horas(inicio, duracao, esperado):
    assert faixa_de_horas([_b("1", 1, inicio, duracao)]) == esperado


def test_a_faixa_cobre_o_bloco_mais_cedo_e_o_mais_tarde():
    blocos = [_b("a", 1, "09:30", 60), _b("b", 5, "21:15", 90)]
    inicio, fim = faixa_de_horas(blocos)
    assert inicio == 9
    assert fim == 23   # 21:15 + 1h30 = 22:45, e a hora cheia seguinte é 23


# ─── A aba da semana ──────────────────────────────────────────────────────


def test_as_duas_abas_nesta_ordem():
    assert _abrir([_b("1", 1, "19:00", 60)]).sheetnames == ["Semana", "Lista"]


def test_a_grade_tem_o_cabecalho_dos_sete_dias():
    ws = _abrir([_b("1", 1, "19:00", 60)])["Semana"]
    cabecalho = [ws.cell(row=4, column=c).value for c in range(1, 9)]
    assert cabecalho == [
        "Hora", "Segunda-feira", "Terça-feira", "Quarta-feira", "Quinta-feira",
        "Sexta-feira", "Sábado", "Domingo",
    ]


def test_o_bloco_cai_na_celula_mesclada_do_dia_e_da_hora_certos():
    """Faixa 17–23, grade a partir da linha 5, 4 linhas por hora.

    Terça (coluna C) 19:00–20:30 → linhas 13 a 18; quinta (coluna E) 19:15–20:45
    → 14 a 19. Os números são a conta, escrita à mão — não a função sob teste.
    """
    ws = _abrir([
        _b("ter", 2, "19:00", 90, "Estequiometria"),
        _b("qui", 4, "19:15", 90, "Revisar meus erros"),
    ])["Semana"]
    assert "C13:C18" in _mesclas(ws)
    assert "E14:E19" in _mesclas(ws)
    assert ws["C13"].value == "Estequiometria\n19:00–20:30"
    assert ws["E14"].value == "Revisar meus erros\n19:15–20:45"


def test_as_horas_cheias_ficam_na_coluna_da_esquerda():
    ws = _abrir([_b("1", 2, "19:00", 90)])["Semana"]
    # 17:00 abre a grade (linha 5) e mescla as quatro linhas da hora.
    assert ws["A5"].value == "17:00"
    assert "A5:A8" in _mesclas(ws)
    assert ws["A13"].value == "19:00"


def test_blocos_encostados_nao_se_sobrepoem_na_grade():
    ws = _abrir([_b("a", 4, "18:00", 60), _b("b", 4, "19:00", 90)])["Semana"]
    assert "E9:E12" in _mesclas(ws)    # 18:00–19:00
    assert "E13:E18" in _mesclas(ws)   # 19:00–20:30


def test_hora_quebrada_nao_gera_mescla_sobreposta():
    """⚠️ A 18:00–19:05 arredonda o fim para 19:15; a 19:10–19:50 arredondaria o
    início para 19:00. Sem o recuo, as duas mesclariam as mesmas linhas — e o Excel
    se recusa a abrir um arquivo assim. A faixa aqui é 16–22 (a folga se divide)."""
    ws = _abrir([_b("a", 4, "18:00", 65, "A"), _b("b", 4, "19:10", 40, "B")])["Semana"]
    faixas = [r for r in ws.merged_cells.ranges if r.min_col == 5]
    assert len(faixas) == 2
    for um, outro in combinations(faixas, 2):
        sobrepoe = um.min_row <= outro.max_row and outro.min_row <= um.max_row
        assert not sobrepoe, f"{um} e {outro} se sobrepõem"
    assert ws["E13"].value.startswith("A\n")   # 18:00 com a grade começando às 16:00


def test_bloco_de_um_quarto_de_hora_ocupa_uma_linha_sem_mescla():
    ws = _abrir([_b("1", 3, "19:00", 15, "Curto")])["Semana"]
    assert ws["D13"].value == "Curto\n19:00–19:15"
    assert not any(r.min_col == 4 for r in ws.merged_cells.ranges)


def test_o_bloco_que_termina_a_meia_noite_cabe_na_grade():
    ws = _abrir([_b("1", 1, "23:00", 60, "Tarde")])["Semana"]   # faixa 18–24
    assert any(r.min_col == 2 for r in ws.merged_cells.ranges)
    assert ws["B25"].value == "Tarde\n23:00–00:00"


def test_o_subtitulo_traz_o_nome_e_a_data():
    ws = _abrir([_b("1", 1, "19:00", 60)], nome="BENNY")["Semana"]
    assert ws["A2"].value == "BENNY · gerado em 24/10/2026"


def test_sem_nome_o_subtitulo_traz_so_a_data():
    ws = _abrir([_b("1", 1, "19:00", 60)], nome=None)["Semana"]
    assert ws["A2"].value == "gerado em 24/10/2026"


def test_a_semana_imprime_em_uma_pagina_deitada():
    """É o horário que se cola na parede. Sem isto o Excel parte a semana em duas
    folhas no meio de quinta-feira."""
    ws = _abrir([_b("1", 1, "19:00", 60)])["Semana"]
    assert ws.page_setup.orientation == "landscape"
    assert ws.page_setup.fitToWidth == 1
    assert ws.page_setup.fitToHeight == 1
    assert ws.sheet_properties.pageSetUpPr.fitToPage is True


# ─── A aba da lista ───────────────────────────────────────────────────────


def _linhas_da_lista(ws) -> list[list]:
    return [[c.value for c in linha] for linha in ws.iter_rows(min_row=1, max_col=5)]


def test_a_lista_vem_por_dia_e_hora_com_a_duracao_em_numero():
    ws = _abrir([
        _b("c", 4, "19:15", 90, "Revisar meus erros"),
        _b("a", 1, "19:00", 90, "Termodinâmica"),
        _b("b", 4, "18:00", 60, "Redação"),
    ])["Lista"]
    assert _linhas_da_lista(ws) == [
        ["Dia", "Início", "Fim", "Duração (min)", "Atividade"],
        ["Segunda-feira", "19:00", "20:30", 90, "Termodinâmica"],
        ["Quinta-feira", "18:00", "19:00", 60, "Redação"],
        ["Quinta-feira", "19:15", "20:45", 90, "Revisar meus erros"],
        ["Total da semana", None, None, 240, None],
    ]
    # NÚMERO: é o que torna a coluna somável no Excel.
    assert isinstance(ws["D2"].value, int)


def test_o_total_fica_fora_do_filtro():
    """Se o filtro cobrisse o total, ordenar a coluna o embaralharia com os blocos."""
    ws = _abrir([_b("a", 1, "19:00", 90), _b("b", 2, "19:00", 60)])["Lista"]
    assert ws.auto_filter.ref == "A1:E3"
    assert ws["A4"].value == "Total da semana"


def test_o_total_soma_a_semana_inteira():
    blocos = [_b(str(i), 1 + i % 7, f"{8 + i:02d}:00", 45) for i in range(8)]
    ws = _abrir(blocos)["Lista"]
    assert ws.cell(row=len(blocos) + 2, column=4).value == 8 * 45


# ─── Texto do aluno nunca vira fórmula ────────────────────────────────────


@pytest.mark.parametrize("rotulo", ["=1+1", "=HYPERLINK(\"http://x\",\"x\")", "+cmd", "-2+3", "@SOMA(A1)"])
def test_rotulo_que_parece_formula_continua_sendo_texto(rotulo):
    """⚠️ Relê o arquivo SALVO. É o `data_type` que o Excel lê: 'f' é fórmula e
    'n' seria número — só 's' garante que a célula abre como o aluno escreveu."""
    wb = _abrir([_b("1", 2, "19:00", 90, rotulo)])
    semana = wb["Semana"]["C13"]
    lista = wb["Lista"]["E2"]
    assert lista.data_type == "s"
    assert lista.value == rotulo
    assert semana.data_type == "s"
    assert semana.value.startswith(rotulo)


def test_o_nome_do_aluno_tambem_nao_vira_formula():
    ws = _abrir([_b("1", 1, "19:00", 60)], nome="=cmd")["Semana"]
    assert ws["A2"].data_type == "s"


def test_sem_blocos_nao_ha_o_que_exportar():
    with pytest.raises(ValueError):
        planilha_do_horario([], "BENNY", HOJE)


# ─── Os dois temas do site ────────────────────────────────────────────────

_CSS = Path(__file__).resolve().parents[2] / "web" / "styles" / "paleta.css"

# O nome na planilha → o token do site.
_TOKENS = {
    "fundo": "fundo", "superficie": "superficie", "superficie_2": "superficie-2", "borda": "borda",
    "acao": "acao", "acao_texto": "acao-texto", "dado": "dado", "dado_claro": "dado-claro",
    "magnitude": "magnitude", "texto": "texto", "texto_2": "texto-2",
}


@pytest.mark.skipif(not _CSS.exists(), reason="web/styles/paleta.css não está neste checkout")
@pytest.mark.parametrize("tema", ["dia", "noite"])
def test_as_paletas_sao_as_do_site(tema):
    """⚠️ O que mantém a planilha com a cara do site quando alguém mexe na paleta."""
    css = _CSS.read_text(encoding="utf-8")
    for campo, token in _TOKENS.items():
        achou = re.search(rf"--{tema}-{re.escape(token)}:\s*#([0-9a-fA-F]{{6}})", css)
        assert achou, f"--{tema}-{token} sumiu de paleta.css"
        assert PALETAS[tema][campo] == achou.group(1).upper(), f"{tema}/{campo} divergiu de paleta.css"


def test_so_existem_os_dois_temas_do_site():
    assert set(PALETAS) == {"dia", "noite"}
    assert set(PALETAS["dia"]) == set(PALETAS["noite"]) == set(_TOKENS)


def test_tema_desconhecido_e_recusado():
    with pytest.raises(ValueError, match="Tema desconhecido"):
        planilha_do_horario([_b("1", 1, "19:00", 60)], "BENNY", HOJE, "roxo")  # type: ignore[arg-type]


def test_sem_tema_e_o_dia():
    ws = _abrir([_b("1", 1, "19:00", 60)])["Semana"]
    assert _cor(ws["A4"]) == PALETAS["dia"]["acao"]


@pytest.mark.parametrize("tema", ["dia", "noite"])
def test_o_cabecalho_usa_a_cor_de_acao_do_tema(tema):
    """Azul-marinho no dia e ouro na noite — a mesma troca do botão do site."""
    p = PALETAS[tema]
    wb = _abrir([_b("1", 2, "19:00", 90)], tema=tema)
    for aba, celula in (("Semana", "B4"), ("Lista", "A1")):
        c = wb[aba][celula]
        assert _cor(c) == p["acao"], f"{aba}!{celula}"
        assert _cor(c.font) == p["acao_texto"], f"{aba}!{celula} (letra)"


def test_a_noite_e_escura_de_verdade_e_o_dia_nao():
    escuro = _abrir([_b("1", 2, "19:00", 90)], tema="noite")["Semana"]
    claro = _abrir([_b("1", 2, "19:00", 90)], tema="dia")["Semana"]
    # Uma célula VAZIA da grade (quarta, 17:00): é a superfície do cartão do site.
    assert _cor(escuro["D5"]) == PALETAS["noite"]["superficie"]
    assert _cor(claro["D5"]) == PALETAS["dia"]["superficie"]
    assert _cor(escuro["D5"]) != _cor(claro["D5"])


def test_o_bloco_sobe_da_superficie_branco_no_dia_e_superficie_2_na_noite():
    """A mesma regra de `dominio/horarioEstudosFolha.ts`."""
    for tema, fundo in (("dia", "fundo"), ("noite", "superficie_2")):
        ws = _abrir([_b("1", 2, "19:00", 90, "Estequiometria")], tema=tema)["Semana"]
        bloco = ws["C13"]
        assert _cor(bloco) == PALETAS[tema][fundo], tema
        assert _cor(bloco.font) == PALETAS[tema]["magnitude"], tema
        # A borda é a cor de DADO — o papel de "o que está planejado".
        assert str(bloco.border.top.color.rgb)[-6:] == PALETAS[tema]["dado"], tema


def test_na_noite_a_coluna_inteira_fica_no_fundo_do_tema():
    """Sem isto só as células escritas ficam escuras, e tudo abaixo da grade abre branco."""
    ws = _abrir([_b("1", 2, "19:00", 90)], tema="noite")["Semana"]
    assert _cor(ws.column_dimensions["B"]) == PALETAS["noite"]["fundo"]
    lista = _abrir([_b("1", 2, "19:00", 90)], tema="noite")["Lista"]
    assert _cor(lista.column_dimensions["A"]) == PALETAS["noite"]["fundo"]


def test_nenhuma_celula_escrita_fica_branca_na_noite():
    """⚠️ O defeito que o Quick Look achou: uma célula com fonte própria não herda o fundo
    da coluna. O título, em letra branca sobre branco, sumia ("o de estudos")."""
    wb = _abrir([_b("1", 2, "19:00", 90, "Estequiometria")], tema="noite")
    fundo = PALETAS["noite"]["fundo"]
    semana, lista = wb["Semana"], wb["Lista"]
    assert _cor(semana["A1"]) == fundo   # o título
    assert _cor(semana["A2"]) == fundo   # o subtítulo
    for linha in lista.iter_rows(min_row=2, max_row=3, max_col=5):
        for c in linha:
            assert _cor(c) == fundo, f"Lista!{c.coordinate} abriria branca no escuro"
    # E nenhuma célula do uso está sem preenchimento: todas têm um fundo do tema. As de
    # DENTRO de uma mescla ficam de fora da conta — o Excel desenha a mescla com o estilo da
    # célula âncora (a do canto), e as outras nunca aparecem.
    for ws in (semana, lista):
        internas = {
            c.coordinate
            for r in ws.merged_cells.ranges
            for linha in ws.iter_rows(min_row=r.min_row, max_row=r.max_row, min_col=r.min_col, max_col=r.max_col)
            for c in linha
            if (c.row, c.column) != (r.min_row, r.min_col)
        }
        sem_fundo = [
            c.coordinate for linha in ws.iter_rows() for c in linha
            if c.fill.fill_type is None and c.coordinate not in internas
        ]
        assert sem_fundo == [], f"{ws.title}: {sem_fundo[:5]}"


def test_a_letra_do_titulo_nunca_fica_igual_ao_fundo():
    """A forma geral do defeito acima, nos dois temas: letra e fundo precisam contrastar."""
    for tema in ("dia", "noite"):
        ws = _abrir([_b("1", 2, "19:00", 90)], tema=tema)["Semana"]
        for coord in ("A1", "A2"):
            assert _cor(ws[coord].font) != _cor(ws[coord]), f"{tema}/{coord}: letra e fundo iguais"


def test_a_letra_da_lista_acompanha_o_tema():
    for tema in ("dia", "noite"):
        ws = _abrir([_b("1", 2, "19:00", 90, "Estequiometria")], tema=tema)["Lista"]
        assert _cor(ws["E2"].font) == PALETAS[tema]["texto"], tema
        assert _cor(ws["A3"].font) == PALETAS[tema]["magnitude"], tema   # "Total da semana"


@pytest.mark.parametrize("tema", ["dia", "noite"])
def test_o_conteudo_e_o_mesmo_nos_dois_temas(tema):
    """O tema muda a cor, e NUNCA o dado: mesma grade, mesma lista, mesmas mesclas."""
    blocos = [_b("a", 2, "19:00", 90, "Estequiometria"), _b("b", 4, "19:15", 90, "Revisar meus erros")]
    de_referencia = _abrir(blocos, tema="dia")
    este = _abrir(blocos, tema=tema)
    assert _mesclas(este["Semana"]) == _mesclas(de_referencia["Semana"])
    assert _linhas_da_lista(este["Lista"]) == _linhas_da_lista(de_referencia["Lista"])
    assert este["Semana"]["C13"].value == de_referencia["Semana"]["C13"].value


def test_o_tema_e_um_literal_de_dia_ou_noite():
    from pydantic import TypeAdapter, ValidationError

    adaptador = TypeAdapter(TemaDaPlanilha)
    assert adaptador.validate_python("noite") == "noite"
    with pytest.raises(ValidationError):
        adaptador.validate_python("roxo")


def test_a_rota_recebe_o_tema_por_query_e_o_valida():
    """O FastAPI valida o `Literal` e responde 422 para qualquer outro valor."""
    from app.routes import horario_estudos as rotas

    rota = next(r for r in rotas.router.routes if r.path.endswith("/planilha"))
    parametros = {p.name: p for p in rota.dependant.query_params}
    assert "tema" in parametros
    assert parametros["tema"].required is False   # sem o parâmetro, é o dia


def test_a_rota_entrega_a_planilha_no_tema_pedido(monkeypatch):
    db = {"bloco_horario_estudo": {"b1": _linha("b1", "aluno-ana", 2, "19:00", 90, "Estequiometria")}}
    rotas = _rota(monkeypatch, db)
    for tema in ("dia", "noite"):
        resposta = asyncio.run(rotas.baixar_planilha(ANA, tema))
        ws = load_workbook(io.BytesIO(resposta.body))["Semana"]
        assert _cor(ws["A4"]) == PALETAS[tema]["acao"], tema


# ─── A rota ───────────────────────────────────────────────────────────────

ANA = {"aluno_id": "aluno-ana", "tipo": "aluno", "nome": "ANA"}


def _rota(monkeypatch, db):
    from app.routes import horario_estudos as rotas

    monkeypatch.setattr(rotas, "get_supabase", lambda: FakeCliente(db))
    return rotas


def _linha(id_, aluno, dia, hora, duracao, rotulo):
    return {
        "id": id_, "aluno_id": aluno, "dia_semana": dia, "hora_inicio": hora,
        "duracao_minutos": duracao, "rotulo": rotulo,
    }


def test_a_rota_entrega_um_xlsx_valido_com_o_nome_certo(monkeypatch):
    db = {"bloco_horario_estudo": {"b1": _linha("b1", "aluno-ana", 2, "19:00:00", 90, "Estequiometria")}}
    resposta = asyncio.run(_rota(monkeypatch, db).baixar_planilha(ANA))

    assert resposta.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    # tema e dia no nome: o claro e o escuro do mesmo dia não se sobrescrevem
    assert re.fullmatch(
        r'attachment; filename="horario-de-estudos-dia-\d{4}-\d{2}-\d{2}\.xlsx"',
        resposta.headers["content-disposition"],
    )
    wb = load_workbook(io.BytesIO(resposta.body))
    # "19:00:00" é o que o PostgREST devolve; a planilha fala "19:00".
    assert wb["Lista"]["B2"].value == "19:00"
    assert wb["Semana"]["A2"].value.startswith("ANA · gerado em ")


def test_a_rota_so_exporta_o_horario_de_quem_pediu(monkeypatch):
    db = {"bloco_horario_estudo": {
        "b1": _linha("b1", "aluno-ana", 2, "19:00", 90, "Da Ana"),
        "b2": _linha("b2", "aluno-bia", 3, "19:00", 90, "Da Bia"),
    }}
    resposta = asyncio.run(_rota(monkeypatch, db).baixar_planilha(ANA))

    lista = load_workbook(io.BytesIO(resposta.body))["Lista"]
    atividades = [lista.cell(row=r, column=5).value for r in range(2, 6)]
    assert "Da Ana" in atividades
    assert "Da Bia" not in atividades


def test_horario_vazio_e_409_e_nao_uma_planilha_em_branco(monkeypatch):
    rotas = _rota(monkeypatch, {"bloco_horario_estudo": {}})
    with pytest.raises(HTTPException) as erro:
        asyncio.run(rotas.baixar_planilha(ANA))
    assert erro.value.status_code == 409
    assert "vazio" in erro.value.detail


def test_o_nome_do_arquivo_e_ascii_mesmo_com_acento_no_nome_do_aluno(monkeypatch):
    """Vai num cabeçalho HTTP: um acento do nome do aluno o quebraria."""
    db = {"bloco_horario_estudo": {"b1": _linha("b1", "aluno-jo", 1, "19:00", 60, "x")}}
    jose = {"aluno_id": "aluno-jo", "tipo": "aluno", "nome": "JOSÉ AÇÃO"}
    resposta = asyncio.run(_rota(monkeypatch, db).baixar_planilha(jose))
    resposta.headers["content-disposition"].encode("ascii")   # não levanta


def test_o_claro_e_o_escuro_do_mesmo_dia_tem_nomes_diferentes(monkeypatch):
    """⚠️ Foi o defeito que a verificação no browser achou: os dois se chamavam
    `horario-de-estudos.xlsx` e o segundo download sobrescrevia o primeiro."""
    db = {"bloco_horario_estudo": {"b1": _linha("b1", "aluno-ana", 2, "19:00", 90, "x")}}
    rotas = _rota(monkeypatch, db)
    nome = {
        tema: asyncio.run(rotas.baixar_planilha(ANA, tema)).headers["content-disposition"]
        for tema in ("dia", "noite")
    }
    assert nome["dia"] != nome["noite"]
    assert "-dia-" in nome["dia"]
    assert "-noite-" in nome["noite"]
