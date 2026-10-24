"""A planilha do horário de estudos (docs/42): a semana do aluno em Excel.

Duas abas, porque servem a coisas diferentes:

  **Semana**  a grade — blocos como células mescladas, dias nas colunas e o tempo
              nas linhas, de 15 em 15 minutos. É o horário que se imprime e se
              cola na parede, e já sai configurado para caber numa página.
  **Lista**   uma linha por bloco, com a duração em minutos como NÚMERO. É a aba
              que se ordena, se filtra e se soma — a grade não deixa.

Sai do backend, e não do navegador, pelo mesmo motivo de toda planilha do
projeto: XLSX não se faz no navegador sem uma biblioteca de ~1 MB
(`cantina_relatorio.py`). PDF e PNG saem do navegador, e a planilha pede a rota
autenticada — dado do próprio aluno, com o `aluno_id` do token.

⚠️ **Os dois temas do site, à escolha de quem exporta** (`tema="dia"` ou `"noite"`).
As cores são as de `web/styles/paleta.css` — `PALETAS` abaixo — e
`test_horario_estudos_planilha.py` LÊ aquele arquivo e falha se uma divergir. É o que
mantém a planilha com a cara do site quando alguém mexe na paleta. O cabeçalho usa a
cor de AÇÃO: azul-marinho no dia e ouro na noite, como o botão do site.

⚠️ **Texto do aluno nunca vira fórmula.** O rótulo é texto livre, e o openpyxl
grava como FÓRMULA toda string que começa com `=`. Um bloco chamado "=1+1" abriria
no Excel como `2`, e um bloco malicioso num arquivo repassado ao grupo da turma
viraria célula executável na máquina de quem abriu. `_texto` força o tipo texto.

⚠️ **Não usa `Workbook(write_only=True)`**, ao contrário de `cantina_relatorio._xlsx`:
o modo de escrita contínua não tem célula mesclada nem estilo por célula, e a
grade é feita deles. O custo — a planilha inteira em memória — é de no máximo 60
blocos.
"""

from __future__ import annotations

import io
import math
from datetime import date
from typing import Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.worksheet import Worksheet

from .horario_estudos import NOME_DO_DIA, hora_de, minutos_de
from .schemas.horario_estudos import BlocoDeEstudo

#: Cada linha da grade é um quarto de hora. 15 minutos é o menor bloco que a API
#: aceita, então nenhum bloco some por ser mais curto que uma linha.
MINUTOS_NO_DIA = 24 * 60
PASSO_EM_MINUTOS = 15
LINHAS_POR_HORA = 60 // PASSO_EM_MINUTOS
#: A grade nunca fica mais curta que isto: um horário só de 19h às 20h30 virando
#: duas linhas de altura seria uma faixa, não um horário.
MINIMO_DE_HORAS = 6

TemaDaPlanilha = Literal["dia", "noite"]

#: ⚠️ Os valores de `web/styles/paleta.css` (`--dia-*` e `--noite-*`), sem o `#`, que é
#: como o openpyxl os quer. Um teste lê aquele arquivo e confere cada um.
PALETAS: dict[str, dict[str, str]] = {
    "dia": {
        "fundo": "FFFFFF", "superficie": "F4F7FC", "superficie_2": "EAF0FA", "borda": "DCE6F7",
        "acao": "1B3F8B", "acao_texto": "FFFFFF", "dado": "2E6BE6", "dado_claro": "7FB6FF",
        "magnitude": "0F1B33", "texto": "16233D", "texto_2": "5C6883",
    },
    "noite": {
        "fundo": "050A18", "superficie": "0C1530", "superficie_2": "101C3D", "borda": "1B2B57",
        "acao": "FFCE3A", "acao_texto": "050A18", "dado": "2F6BFF", "dado_claro": "7FB6FF",
        "magnitude": "FFFFFF", "texto": "E6EDFB", "texto_2": "6E85B8",
    },
}

#: Começos que o Excel interpreta como FÓRMULA. Só o `=` vira fórmula no openpyxl,
#: mas a lista é a mesma de `cantina_relatorio._INICIO_DE_FORMULA`: o cinto vale
#: para qualquer um deles.
_INICIO_DE_FORMULA = ("=", "+", "-", "@", "\t", "\r")

_COLUNA_DO_PRIMEIRO_DIA = 2   # B; a A é a das horas
_PRIMEIRA_LINHA_DA_GRADE = 5  # 1 título, 2 subtítulo, 3 vazia, 4 cabeçalho


def faixa_de_horas(blocos: list[BlocoDeEstudo]) -> tuple[int, int]:
    """(hora inicial, hora final) da grade — horas inteiras, no máximo 0..24.

    O início sai da hora cheia do bloco mais cedo, e o fim da hora cheia seguinte
    ao bloco mais tarde. Se sobrar pouco, a folga se DIVIDE dos dois lados — quem
    estuda só às 19h vê 17h–23h, e não 18h–24h com um terço da página vazio no fim
    — e, batendo na meia-noite (ou no zero), o que sobrou vai para o outro lado.

    ⚠️ É a MESMA regra de `dominio/horarioEstudosFolha.ts`, e a repetição é
    deliberada: a grade do PDF e a do Excel têm de mostrar o mesmo recorte do dia.
    """
    inicio = min(minutos_de(b.horaInicio) for b in blocos) // 60
    fim = math.ceil(max(minutos_de(b.horaInicio) + b.duracaoMinutos for b in blocos) / 60)
    falta = MINIMO_DE_HORAS - (fim - inicio)
    if falta > 0:
        novo_inicio = max(0, inicio - falta // 2)
        novo_fim = min(24, fim + (falta - (inicio - novo_inicio)))
        sobrou = MINIMO_DE_HORAS - (novo_fim - novo_inicio)
        if sobrou > 0:
            novo_inicio = max(0, novo_inicio - sobrou)
        inicio, fim = novo_inicio, novo_fim
    return inicio, fim


def _fim(bloco: BlocoDeEstudo) -> str:
    """"20:30". O bloco que termina à meia-noite diz "00:00", como a tela — e não o
    "24:00" de `hora_de`, que o aluno nunca vê escrito em lugar nenhum do app."""
    return hora_de((minutos_de(bloco.horaInicio) + bloco.duracaoMinutos) % MINUTOS_NO_DIA)


def _texto(celula, valor: str) -> None:
    """Grava `valor` como TEXTO, mesmo que comece como uma fórmula."""
    celula.value = valor
    if valor.startswith(_INICIO_DE_FORMULA):
        celula.data_type = "s"


def _fino(cor: str) -> Side:
    return Side(style="thin", color=cor)


def _preenchimento(cor: str) -> PatternFill:
    return PatternFill("solid", fgColor=cor)


def _cabecalho(ws: Worksheet, linha: int, titulos: list[str], p: dict[str, str]) -> None:
    for i, titulo in enumerate(titulos, start=1):
        c = ws.cell(row=linha, column=i, value=titulo)
        c.font = Font(bold=True, color=p["acao_texto"])
        c.fill = _preenchimento(p["acao"])
        c.alignment = Alignment(horizontal="center", vertical="center")


def _colunas_no_fundo(ws: Worksheet, letras: list[str], p: dict[str, str]) -> None:
    """A coluna INTEIRA no fundo do tema.

    Sem isto, no escuro só as células escritas ficam escuras, e tudo abaixo da grade
    abre branco no meio de uma planilha preta. Cobre as células VAZIAS; as que têm
    estilo próprio precisam de `_fundo_no_que_sobrou`.
    """
    for letra in letras:
        ws.column_dimensions[letra].fill = _preenchimento(p["fundo"])


def _fundo_no_que_sobrou(ws: Worksheet, p: dict[str, str]) -> None:
    """Dá o fundo do tema a toda célula escrita que ficou sem preenchimento.

    ⚠️ O estilo da COLUNA só vale para a célula que não tem estilo nenhum. Basta uma
    fonte (o título, o subtítulo, cada linha da lista) para a célula deixar de herdá-lo
    e abrir BRANCA — e o título da planilha da noite, em letra branca, sumia. Achado
    pelo Quick Look, na verificação: "Horário de estudos" saía como "o de estudos".
    """
    fundo = _preenchimento(p["fundo"])
    for linha in ws.iter_rows():
        for c in linha:
            if c.fill.fill_type is None:
                c.fill = fundo


def _aba_da_semana(
    wb: Workbook, blocos: list[BlocoDeEstudo], nome: str | None, gerado_em: date,
    tema: str, p: dict[str, str],
) -> None:
    ws = wb.active
    ws.title = "Semana"
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = p["dado"]
    inicio_h, fim_h = faixa_de_horas(blocos)
    horas = fim_h - inicio_h
    total_de_linhas = horas * LINHAS_POR_HORA
    primeira = _PRIMEIRA_LINHA_DA_GRADE
    ultima_coluna = _COLUNA_DO_PRIMEIRO_DIA + len(NOME_DO_DIA) - 1
    _colunas_no_fundo(ws, [get_column_letter(c) for c in range(1, ultima_coluna + 1)], p)

    ws.cell(row=1, column=1, value="Horário de estudos").font = Font(bold=True, size=16, color=p["magnitude"])
    sub = ws.cell(row=2, column=1)
    _texto(sub, " · ".join(parte for parte in (nome, f"gerado em {gerado_em:%d/%m/%Y}") if parte))
    sub.font = Font(color=p["texto_2"])

    _cabecalho(ws, primeira - 1, ["Hora", *NOME_DO_DIA.values()], p)

    # A grade vazia primeiro — linha e hora —, e os blocos por cima. A superfície é a do
    # cartão do site, e a borda é a dele; a hora cheia ganha a linha de baixo.
    for deslocamento in range(total_de_linhas):
        linha = primeira + deslocamento
        ws.row_dimensions[linha].height = 15
        fecha_a_hora = (deslocamento + 1) % LINHAS_POR_HORA == 0
        for coluna in range(1, ultima_coluna + 1):
            c = ws.cell(row=linha, column=coluna)
            c.fill = _preenchimento(p["superficie"])
            c.border = Border(
                left=_fino(p["borda"]), right=_fino(p["borda"]),
                bottom=_fino(p["borda"]) if fecha_a_hora else None,
            )
    for h in range(horas):
        topo = primeira + h * LINHAS_POR_HORA
        ws.merge_cells(start_row=topo, end_row=topo + LINHAS_POR_HORA - 1, start_column=1, end_column=1)
        rotulo = ws.cell(row=topo, column=1, value=f"{inicio_h + h:02d}:00")
        rotulo.font = Font(bold=True, color=p["texto_2"], size=9)
        rotulo.alignment = Alignment(horizontal="center", vertical="top")

    # O bloco SOBE da superfície: branco no dia, `superficie-2` na noite — a mesma regra
    # de `dominio/horarioEstudosFolha.ts`. A borda é a cor de DADO (o que está planejado).
    fundo_do_bloco = p["fundo"] if tema == "dia" else p["superficie_2"]
    for dia in NOME_DO_DIA:
        coluna = _COLUNA_DO_PRIMEIRO_DIA + dia - 1
        do_dia = sorted(
            (b for b in blocos if b.diaSemana == dia), key=lambda b: minutos_de(b.horaInicio)
        )
        livre = 0   # primeira linha ainda sem bloco, para dois blocos nunca se mesclarem sobre a mesma célula
        for b in do_dia:
            inicio = minutos_de(b.horaInicio)
            fim = inicio + b.duracaoMinutos
            ini_linha = max(inicio // PASSO_EM_MINUTOS - inicio_h * LINHAS_POR_HORA, livre)
            fim_linha = max(math.ceil(fim / PASSO_EM_MINUTOS) - inicio_h * LINHAS_POR_HORA, ini_linha + 1)
            livre = fim_linha
            topo, base = primeira + ini_linha, primeira + fim_linha - 1
            if base > topo:
                ws.merge_cells(start_row=topo, end_row=base, start_column=coluna, end_column=coluna)
            for linha in range(topo, base + 1):
                c = ws.cell(row=linha, column=coluna)
                c.fill = _preenchimento(fundo_do_bloco)
                c.border = Border(
                    left=_fino(p["dado"]), right=_fino(p["dado"]),
                    top=_fino(p["dado"]) if linha == topo else None,
                    bottom=_fino(p["dado"]) if linha == base else None,
                )
            celula = ws.cell(row=topo, column=coluna)
            _texto(celula, f"{b.rotulo}\n{b.horaInicio}–{_fim(b)}")
            celula.font = Font(bold=True, color=p["magnitude"], size=10)
            celula.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")

    ws.column_dimensions["A"].width = 8
    for dia in NOME_DO_DIA:
        ws.column_dimensions[chr(ord("A") + dia)].width = 24
    ws.freeze_panes = ws.cell(row=primeira, column=_COLUNA_DO_PRIMEIRO_DIA)

    # Cabe numa página, deitada: é o horário que se imprime. Sem isto o Excel parte a
    # semana em duas folhas no meio de quinta-feira.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    _fundo_no_que_sobrou(ws, p)


def _aba_da_lista(wb: Workbook, blocos: list[BlocoDeEstudo], p: dict[str, str]) -> None:
    ws = wb.create_sheet("Lista")
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = p["dado"]
    _colunas_no_fundo(ws, list("ABCDE"), p)
    _cabecalho(ws, 1, ["Dia", "Início", "Fim", "Duração (min)", "Atividade"], p)

    ordenados = sorted(blocos, key=lambda b: (b.diaSemana, minutos_de(b.horaInicio)))
    for i, b in enumerate(ordenados, start=2):
        ws.cell(row=i, column=1, value=NOME_DO_DIA[b.diaSemana])
        ws.cell(row=i, column=2, value=b.horaInicio)
        ws.cell(row=i, column=3, value=_fim(b))
        # NÚMERO, não texto: é o que torna a coluna somável.
        ws.cell(row=i, column=4, value=b.duracaoMinutos)
        _texto(ws.cell(row=i, column=5), b.rotulo)
        for coluna in range(1, 6):
            c = ws.cell(row=i, column=coluna)
            c.font = Font(color=p["texto"])
            c.border = Border(bottom=_fino(p["borda"]))
            if coluna in (2, 3):
                c.alignment = Alignment(horizontal="center")

    total = len(ordenados) + 2
    for coluna in range(1, 6):
        ws.cell(row=total, column=coluna).border = Border(top=_fino(p["dado"]))
    rotulo = ws.cell(row=total, column=1, value="Total da semana")
    rotulo.font = Font(bold=True, color=p["magnitude"])
    soma = ws.cell(row=total, column=4, value=sum(b.duracaoMinutos for b in ordenados))
    soma.font = Font(bold=True, color=p["magnitude"])

    for letra, largura in zip("ABCDE", (18, 9, 9, 15, 48), strict=True):
        ws.column_dimensions[letra].width = largura
    ws.freeze_panes = "A2"
    # O filtro cobre só os blocos: o total fica de fora, senão ordenar o embaralharia.
    ws.auto_filter.ref = f"A1:E{len(ordenados) + 1}"
    _fundo_no_que_sobrou(ws, p)


def planilha_do_horario(
    blocos: list[BlocoDeEstudo],
    nome: str | None,
    gerado_em: date,
    tema: TemaDaPlanilha = "dia",
) -> bytes:
    """O .xlsx do horário, no tema `dia` ou `noite`. `blocos` não pode ser vazio —
    quem chama recusa antes."""
    if not blocos:
        raise ValueError("Não há blocos para exportar.")
    if tema not in PALETAS:
        raise ValueError(f"Tema desconhecido: {tema!r}. Use 'dia' ou 'noite'.")
    paleta = PALETAS[tema]
    wb = Workbook()
    _aba_da_semana(wb, blocos, nome, gerado_em, tema, paleta)
    _aba_da_lista(wb, blocos, paleta)
    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()
