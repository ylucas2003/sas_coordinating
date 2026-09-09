"""O relatório de custos da cantina, em XLSX com gráfico (docs/40 §12.11).

⚠️ **Esta é a primeira rota de exportação da API, e a exceção é consciente.**
O docs/38 §8.2 registrou "não existe rota de exportação — o export é só do lado
do cliente" como escolha; ela muda **aqui e só aqui**, por uma razão que não
tem contorno: **gráfico em XLSX não se faz no navegador.** Nem SheetJS nem
exceljs escrevem gráfico, e a biblioteca ainda pesaria ~1 MB no bundle que a
§12.1.4 acabou de cortar. O `openpyxl` já estava no `requirements.txt` — hoje
só para LER o XLSX do Canvas — e escreve gráfico nativo do Excel.

O CSV e o PDF de hoje continuam no cliente: não se migra o que já funciona.

⚠️ **Nada aqui é cobrança.** Não há fatura, "quem pagou" nem conciliação — a
fronteira do docs/38 §8.1.5 segue de pé. O recorte por aluno é leitura: é o que
mais se aproxima dela, e no dia em que alguém pedir "manda a conta do aluno", a
resposta deixa de ser um relatório e vira outro produto.

⚠️ **A restrição alimentar não entra.** Mesma régua da tela da coordenação
(docs/38 §2.6): quem exporta custo não precisa de dado de saúde de menor.
"""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from .supabase_client import ClienteDados

#: Rótulo das duas refeições, para o cabeçalho não sair em `snake_case`.
_ROTULO = {"almoco": "Almoço", "janta": "Janta"}

#: O modo, na palavra que a cantina usa (docs/40 §8).
_MODO = {"pedido": "pedido", "presencial": "retirada na hora"}


def montar(
    cliente: ClienteDados, *, de: date, ate: date, cantina_id: str | None = None
) -> bytes:
    """O arquivo inteiro, em memória.

    ⚠️ **Chame por `asyncio.to_thread`.** Isto é trabalho de CPU, síncrono, e a
    API tem um event loop só (`UVICORN_WORKERS=1`): montar 40 mil linhas aqui
    dentro de um `async def` congelaria o produto para todo mundo — exatamente
    o defeito que a §12.1.1 acabou de consertar no chat. Seria consertar o
    congelamento e trazê-lo de volta pela porta que este arquivo abriu.
    """
    dados = _ler(cliente, de=de, ate=ate, cantina_id=cantina_id)

    # `write_only`: as linhas são escritas e soltas, em vez de guardadas numa
    # árvore de células. Com um mês de 900 alunos em duas refeições são ~40 mil
    # linhas na aba de pedidos — e o modo normal as manteria todas em memória
    # ao mesmo tempo. Gráfico funciona nos dois modos (conferido).
    wb = Workbook(write_only=True)

    _aba_pedidos(wb, dados)
    _aba_por_dia(wb, dados)
    _aba_por_turma(wb, dados)
    _aba_por_aluno(wb, dados)
    if len(dados["cantinas"]) > 1:
        # Só com mais de uma: uma aba de uma linha só é ruído, e ensina a
        # ignorar as abas do arquivo.
        _aba_por_cantina(wb, dados)
    _aba_cardapio(wb, dados)

    saida = io.BytesIO()
    wb.save(saida)
    return saida.getvalue()


def agregados(
    cliente: ClienteDados, *, de: date, ate: date, cantina_id: str | None = None
) -> dict:
    """Os mesmos números do XLSX, em JSON — para a TELA.

    ⚠️ A tela e a planilha somam pelo MESMO caminho de propósito. Dois
    agregadores para a mesma pergunta divergem no primeiro caso de borda (um
    pedido sem valor, uma turma vazia), e a divergência aparece como "a
    planilha não bate com a tela" — que é a forma mais cara de descobrir que
    havia dois códigos.
    """
    dados = _ler(cliente, de=de, ate=ate, cantina_id=cantina_id)

    def linhas(chave) -> list[dict]:
        return [
            {"rotulo": r, "refeicoes": q, "total": v} for r, q, v in _somar(dados, chave)
        ]

    def de_qual_cantina(p: dict) -> str:
        cardapio = dados["por_cardapio"].get(p["cardapio_id"], {})
        return dados["nome_da_cantina"].get(cardapio.get("cantina_id"), "sem cantina")

    return {
        "de": de.isoformat(),
        "ate": ate.isoformat(),
        "refeicoes": len(dados["pedidos"]),
        "total": round(sum(float(p.get("valor_cobrado") or 0) for p in dados["pedidos"]), 2),
        # ⚠️ Quantos pedidos ficaram SEM preço. O relatório que soma zero em
        # silêncio faz a coordenação fechar a conta errada e nunca saber —
        # dizer "12 sem valor registrado" é o que transforma o buraco em
        # pergunta (docs/40 §12.11.2).
        "semValor": sum(1 for p in dados["pedidos"] if p.get("valor_cobrado") is None),
        "porDia": linhas(
            lambda p: f"{p.get('data')} · {_ROTULO.get(p.get('refeicao'), '')}"
        ),
        "porTurma": linhas(
            lambda p: dados["turma_do_aluno"].get(p["aluno_id"]) or "sem turma"
        ),
        "porAluno": linhas(
            lambda p: dados["nome_do_aluno"].get(p["aluno_id"]) or "sem nome"
        ),
        "porCantina": linhas(de_qual_cantina),
    }


def _ler(
    cliente: ClienteDados, *, de: date, ate: date, cantina_id: str | None
) -> dict:
    """Todas as consultas do relatório, em LOTE.

    Uma consulta por tabela, e nenhuma dentro de laço: é o mesmo cuidado da
    §12.1.3, e aqui ele pesa mais — este é o único lugar do produto que lê um
    mês inteiro de pedidos de uma vez.
    """
    cardapios = (
        cliente.table("cardapio")
        .select("id, cantina_id, data, refeicao")
        .gte("data", de.isoformat())
        .lte("data", ate.isoformat())
        .execute()
        .data
        or []
    )
    if cantina_id:
        cardapios = [c for c in cardapios if c["cantina_id"] == cantina_id]
    ids = [c["id"] for c in cardapios]
    por_cardapio = {c["id"]: c for c in cardapios}

    pedidos = (
        (
            cliente.table("pedido_refeicao")
            .select(
                "id, cardapio_id, aluno_id, modo, valor_cobrado, criado_em, retirado_em, "
                "data, refeicao"
            )
            .in_("cardapio_id", ids)
            .execute()
            .data
            or []
        )
        if ids
        else []
    )

    aluno_ids = sorted({p["aluno_id"] for p in pedidos})
    alunos = (
        (
            cliente.table("aluno")
            .select("id, nome")
            .in_("id", aluno_ids)
            .execute()
            .data
            or []
        )
        if aluno_ids
        else []
    )
    matriculas = (
        (
            cliente.table("matricula_turma")
            .select("aluno_id, turma(section_original)")
            .in_("aluno_id", aluno_ids)
            .is_("ativo_ate", "null")
            .execute()
            .data
            or []
        )
        if aluno_ids
        else []
    )

    cantinas = (
        cliente.table("cantina").select("id, nome").execute().data or []
    )
    usadas = {c["cantina_id"] for c in cardapios}

    blocos = (
        (
            cliente.table("cardapio_bloco")
            .select("id, cardapio_id, nome, ordem, escolhas_minimas, escolhas_maximas, observacao")
            .in_("cardapio_id", ids)
            .execute()
            .data
            or []
        )
        if ids
        else []
    )
    opcoes = (
        (
            cliente.table("cardapio_opcao")
            .select("id, bloco_id, nome, ordem")
            .in_("bloco_id", [b["id"] for b in blocos])
            .execute()
            .data
            or []
        )
        if blocos
        else []
    )
    itens = (
        (
            cliente.table("pedido_refeicao_item")
            .select("pedido_id, opcao_id")
            .in_("pedido_id", [p["id"] for p in pedidos])
            .execute()
            .data
            or []
        )
        if pedidos
        else []
    )

    return {
        "de": de,
        "ate": ate,
        "cardapios": cardapios,
        "por_cardapio": por_cardapio,
        "pedidos": sorted(pedidos, key=lambda p: (p.get("data") or "", p.get("refeicao") or "")),
        "nome_do_aluno": {a["id"]: a["nome"] for a in alunos},
        "turma_do_aluno": {
            m["aluno_id"]: (m.get("turma") or {}).get("section_original") for m in matriculas
        },
        "nome_da_cantina": {c["id"]: c["nome"] for c in cantinas},
        "cantinas": [c for c in cantinas if c["id"] in usadas],
        "blocos": blocos,
        "opcoes": opcoes,
        "itens": itens,
    }


def _bloco_de_cada_opcao(dados: dict) -> tuple[dict[str, str], dict[str, str], list[str]]:
    """`opcao_id → (nome da opção, nome do bloco)`, e a ordem dos blocos.

    A ordem é a da bandeja — `bloco.ordem` —, e não alfabética: a cantina monta
    o cardápio na ordem em que o prato é servido.
    """
    nome_do_bloco = {b["id"]: b["nome"] for b in dados["blocos"]}
    ordem_do_bloco: dict[str, int] = {}
    for b in dados["blocos"]:
        atual = ordem_do_bloco.get(b["nome"])
        if atual is None or b["ordem"] < atual:
            ordem_do_bloco[b["nome"]] = b["ordem"]

    opcao_nome = {o["id"]: o["nome"] for o in dados["opcoes"]}
    opcao_bloco = {o["id"]: nome_do_bloco.get(o["bloco_id"], "") for o in dados["opcoes"]}
    colunas = sorted(ordem_do_bloco, key=lambda nome: ordem_do_bloco[nome])
    return opcao_nome, opcao_bloco, colunas


def _aba_pedidos(wb: Workbook, dados: dict) -> None:
    """Uma linha por aluno-dia, com **uma coluna por bloco** (docs/40 §12.5.2).

    É o formato do formulário que a coordenação já usa, e é o que torna a
    coluna somável: "Arroz | Frango | Folhas" numa célula só não responde
    "quantos pediram frango" sem alguém separar à mão.
    """
    ws = wb.create_sheet("Pedidos")
    opcao_nome, opcao_bloco, colunas = _bloco_de_cada_opcao(dados)

    escolhas_do_pedido: dict[str, dict[str, list[str]]] = defaultdict(lambda: defaultdict(list))
    for item in dados["itens"]:
        bloco = opcao_bloco.get(item["opcao_id"])
        if bloco:
            escolhas_do_pedido[item["pedido_id"]][bloco].append(
                opcao_nome.get(item["opcao_id"], "")
            )

    ws.append(
        ["Data", "Cantina", "Refeição", "Aluno", "Turma", "Modo", *colunas, "Valor", "Hora"]
    )
    for p in dados["pedidos"]:
        cardapio = dados["por_cardapio"].get(p["cardapio_id"], {})
        presencial = (p.get("modo") or "pedido") == "presencial"
        # ⚠️ Traço, e não célula vazia, para quem pega na hora: vazio numa
        # planilha lê-se como "faltou o dado"; o traço diz que não havia dado a
        # ter (docs/40 §10.1).
        celulas = [
            "—" if presencial else "; ".join(escolhas_do_pedido[p["id"]].get(bloco, []))
            for bloco in colunas
        ]
        ws.append([
            p.get("data") or cardapio.get("data"),
            dados["nome_da_cantina"].get(cardapio.get("cantina_id"), ""),
            _ROTULO.get(p.get("refeicao") or cardapio.get("refeicao"), ""),
            dados["nome_do_aluno"].get(p["aluno_id"], ""),
            dados["turma_do_aluno"].get(p["aluno_id"], ""),
            _MODO.get(p.get("modo") or "pedido", ""),
            *celulas,
            p.get("valor_cobrado"),
            (p.get("retirado_em") or p.get("criado_em") or "")[:19].replace("T", " "),
        ])


def _somar(dados: dict, chave) -> list[tuple[str, int, float]]:
    """`(rótulo, quantas refeições, quanto somou)`, ordenado pelo rótulo.

    ⚠️ Soma `valor_cobrado`, que é o preço CONGELADO no instante do pedido
    (migration 0054) — e não o valor de tabela de hoje. É a diferença entre o
    relatório de março continuar dizendo o que março custou e ele mudar quando
    o preço subir em abril.
    """
    contagem: dict[str, int] = defaultdict(int)
    total: dict[str, float] = defaultdict(float)
    for p in dados["pedidos"]:
        rotulo = chave(p)
        if rotulo is None:
            continue
        contagem[rotulo] += 1
        total[rotulo] += float(p.get("valor_cobrado") or 0)
    return [(r, contagem[r], round(total[r], 2)) for r in sorted(contagem)]


def _aba_com_grafico(
    wb: Workbook, titulo: str, cabecalho: str, linhas: list[tuple[str, int, float]]
) -> None:
    ws = wb.create_sheet(titulo)
    ws.append([cabecalho, "Refeições", "Total (R$)"])
    for linha in linhas:
        ws.append(list(linha))

    if not linhas:
        # Sem dado, sem gráfico: um gráfico vazio parece defeito de geração.
        return

    grafico = BarChart()
    grafico.title = titulo
    grafico.y_axis.title = "R$"
    grafico.add_data(
        Reference(ws, min_col=3, min_row=1, max_row=len(linhas) + 1), titles_from_data=True
    )
    grafico.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(linhas) + 1))
    ws.add_chart(grafico, "F2")


def _aba_por_dia(wb: Workbook, dados: dict) -> None:
    _aba_com_grafico(
        wb, "Custos · por dia", "Dia",
        _somar(dados, lambda p: f"{p.get('data')} · {_ROTULO.get(p.get('refeicao'), '')}"),
    )


def _aba_por_turma(wb: Workbook, dados: dict) -> None:
    _aba_com_grafico(
        wb, "Custos · por turma", "Turma",
        _somar(dados, lambda p: dados["turma_do_aluno"].get(p["aluno_id"]) or "sem turma"),
    )


def _aba_por_aluno(wb: Workbook, dados: dict) -> None:
    _aba_com_grafico(
        wb, "Custos · por aluno", "Aluno",
        _somar(dados, lambda p: dados["nome_do_aluno"].get(p["aluno_id"]) or "sem nome"),
    )


def _aba_por_cantina(wb: Workbook, dados: dict) -> None:
    def de_qual(p: dict) -> str:
        cardapio = dados["por_cardapio"].get(p["cardapio_id"], {})
        return dados["nome_da_cantina"].get(cardapio.get("cantina_id"), "sem cantina")

    _aba_com_grafico(wb, "Custos · por cantina", "Cantina", _somar(dados, de_qual))


def _aba_cardapio(wb: Workbook, dados: dict) -> None:
    """A grade semanal: blocos nas linhas, dias nas colunas (docs/40 §12.5.3).

    É a tabela que a cozinha já desenha no papel, incluindo a coluna "Obs!" —
    que aqui junta o que o SAS sabe em número (`escolhas_minimas`/`maximas`)
    com o texto livre da migration 0053.
    """
    ws = wb.create_sheet("Cardápio")

    dias = sorted({(c["data"], c["refeicao"]) for c in dados["cardapios"]})
    ws.append(["Bloco", "Opção", *[f"{d} · {_ROTULO.get(r, '')}" for d, r in dias], "Obs!"])

    opcoes_do_bloco: dict[str, list[dict]] = defaultdict(list)
    for o in dados["opcoes"]:
        opcoes_do_bloco[o["bloco_id"]].append(o)

    por_dia_e_nome: dict[tuple[str, str, str], list[str]] = defaultdict(list)
    obs_do_nome: dict[str, str] = {}
    for b in dados["blocos"]:
        cardapio = dados["por_cardapio"].get(b["cardapio_id"], {})
        chave_dia = (cardapio.get("data"), cardapio.get("refeicao"))
        for o in sorted(opcoes_do_bloco.get(b["id"], []), key=lambda x: x["ordem"]):
            por_dia_e_nome[(b["nome"], chave_dia[0], chave_dia[1])].append(o["nome"])
        limite = _limite_em_palavras(b)
        texto = (b.get("observacao") or "").strip()
        obs_do_nome[b["nome"]] = " · ".join(x for x in (limite, texto) if x)

    for nome_do_bloco in sorted(
        {b["nome"] for b in dados["blocos"]},
        key=lambda n: min(b["ordem"] for b in dados["blocos"] if b["nome"] == n),
    ):
        # Quantas linhas este bloco ocupa: a maior lista de opções entre os dias.
        maior = max(
            (len(por_dia_e_nome.get((nome_do_bloco, d, r), [])) for d, r in dias), default=0
        )
        for i in range(maior):
            ws.append([
                nome_do_bloco if i == 0 else "",
                f"Opção {i + 1}",
                *[
                    (por_dia_e_nome.get((nome_do_bloco, d, r), []) + [""] * maior)[i]
                    for d, r in dias
                ],
                obs_do_nome.get(nome_do_bloco, "") if i == 0 else "",
            ])


def _limite_em_palavras(bloco: dict) -> str:
    """"máximo 2 opções" — o que o par mín./máx. diz, na voz da tabela de papel."""
    minimo = bloco.get("escolhas_minimas") or 0
    maximo = bloco.get("escolhas_maximas") or 0
    if maximo == 0:
        return "só para conferir"
    if minimo == maximo:
        return f"escolha {minimo}"
    if minimo == 0:
        return f"máximo {maximo} {'opção' if maximo == 1 else 'opções'}"
    return f"de {minimo} a {maximo} opções"
